#!/usr/bin/env python3
"""ŠABLONA PODKLADŮ OD CONTROLLINGU (C3 #537, 27. 10. 2026).

Vyrobí podklady/controlling/SABLONA_PODKLADU_CONTROLLING.xlsx — sešit, který
J. V. pošle controllingu: přesně jaká data a v jaké struktuře potřebujeme,
aby šlo dokončit roadmapu #500 (budoucí účty Helios, rozvržení PROJ, sazby
a hodiny PROJ, období, formát čísel, převodník CN ↔ projekt).

Struktura (id a názvy skupin, klíče ceníku PROJ a jejich navržené zařazení)
se čte z KÓDU aplikace (src/controlling.js, src/cenik.js) přes Node — žádná
částka, žádný ceník, žádná osobní data. Každý list s daty má jeden řádek
„VZOR" se SMYŠLENÝMI hodnotami, který ukazuje formát.

Spuštění z kořene repozitáře:
  python3 nastroje/sablona_podkladu_controlling.py [výstup.xlsx]
Kontrola: src/test_sablona_controlling.js (klíče PROJ a skupiny v uloženém
sešitu = kód; když přibude klíč, test spadne → šablonu vyrobit znovu).
"""
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

KOREN = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
VYSTUP = os.path.join(KOREN, 'podklady', 'controlling', 'SABLONA_PODKLADU_CONTROLLING.xlsx')

NODE = r"""
const nacti = f => { const m = require('./src/' + f); Object.keys(m).forEach(k => { if (global[k] === undefined) global[k] = m[k]; }); return m; };
['cenik.js', 'cenik_stari.js', 'engine.js', 'engine_proj.js'].forEach(nacti);
const C = nacti('controlling.js');
const v = C.ctrlVychozi();
const proj = [];
CENIK_DEF_PROJ.forEach(([sek, rows]) => rows.forEach(r => proj.push({ sekce: sek, klic: r[0], popis: r[1], jednotka: r[2], skupina: v.mapa[r[0]] || '' })));
console.log(JSON.stringify({ verzeStruktury: C.CTRL_EXPORT_VERZE, skupiny: v.skupiny, proj,
  typy: C.CTRL_TYP_TEXT || { primy: 'přímý', neprimy: 'nepřímý' } }));
"""

TYP = {'primy': 'přímý', 'neprimy': 'nepřímý'}
CHOVANI = {'variabilni': 'variabilní', 'fixni': 'fixní'}
OBLAST = {'ock': 'OCK', 'proj': 'PROJ', 'obe': 'obě'}

HLAVA = Font(bold=True, color='FFFFFF')
HLAVA_FILL = PatternFill('solid', fgColor='1F4E79')
VZOR_FILL = PatternFill('solid', fgColor='FFF2CC')
KOD_FILL = PatternFill('solid', fgColor='E7EEF7')
ZALOM = Alignment(wrap_text=True, vertical='top')


def data_z_kodu():
    out = subprocess.run(['node', '-e', NODE], cwd=KOREN, capture_output=True, text=True, check=True).stdout
    return json.loads(out)


def list_s_hlavickou(wb, nazev, hlavicka, sirky, popis):
    ws = wb.create_sheet(nazev)
    ws.append([popis])
    ws['A1'].font = Font(italic=True)
    ws['A1'].alignment = ZALOM
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(hlavicka))
    ws.row_dimensions[1].height = 48
    ws.append(hlavicka)
    for i, sirka in enumerate(sirky, start=1):
        c = ws.cell(row=2, column=i)
        c.font, c.fill, c.alignment = HLAVA, HLAVA_FILL, ZALOM
        ws.column_dimensions[c.column_letter].width = sirka
    ws.freeze_panes = 'A3'
    return ws


def obarvi(ws, radek, fill):
    for c in ws[radek]:
        c.fill = fill
        c.alignment = ZALOM


def vyber(ws, sloupec, moznosti, od=3, do=500):
    dv = DataValidation(type='list', formula1='"' + ','.join(moznosti) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{sloupec}{od}:{sloupec}{do}')


def sestav(d):
    wb = Workbook()
    nav = wb.active
    nav.title = 'Návod'
    nav.column_dimensions['A'].width = 28
    nav.column_dimensions['B'].width = 110
    radky = [
        ('Podklady pro controlling v kalkulátoru', ''),
        ('Účel', 'Kalkulátor posílá předkalkulaci (náklady nabídky po nákladových skupinách s účty Helios) do Reporting CZ '
                 'k porovnání se skutečností (PRE × POST). K dokončení potřebujeme od controllingu údaje na listech níže.'),
        ('Co vyplnit', '1) Skupiny — potvrdit číselník, doplnit budoucí účty Helios a případné změny. '
                       '2) Mapování PROJ — potvrdit nebo změnit navržené zařazení položek ceníku projekce. '
                       '3) PROJ sazby a hodiny — nákladové sazby vlastních lidí (role, ne jména) a způsob vykazování hodin. '
                       '4) Převodník — jen ukázat formát, jakým se páruje číslo nabídky s projektem Helios. '
                       '5) Otázky — odpovědi na formát, období a import.'),
        ('Barvy', 'Modré řádky = převzato z kalkulátoru (jen názvy a klíče, bez částek). Žlutý řádek „VZOR" = smyšlený příklad '
                  'formátu, smazat nebo přepsat. Bílé buňky vyplňuje controlling.'),
        ('Čísla', 'Částky v Kč bez DPH jako ČÍSLO (ne text), nejvýš 2 desetinná místa; bez mezer a „Kč" v buňce. '
                  'Procenta jako číslo 0–100. Hodiny jako číslo s desetinnou částí (1,5 = hodina a půl).'),
        ('Data', 'Datum jako datum Excelu nebo text RRRR-MM-DD. Období vždy od–do včetně obou dnů.'),
        ('Účty Helios', 'Šestimístné číslo jako TEXT (např. 518001); více účtů oddělit čárkou. Prázdné = zatím neurčeno.'),
        ('Id skupiny', 'Krátký kód malými písmeny bez diakritiky a mezer (např. sl-lakovani). Existující id neměnit — '
                       'změní se jen název; nová skupina = nové id. Zrušení skupiny napsat do sloupce Změna.'),
        ('Co sem nepatří', 'Jména zaměstnanců, mzdy jednotlivců, rodná čísla, kontakty — sazby se uvádějí za roli. '
                           'Skutečné náklady zakázek zůstávají v Reporting CZ, do kalkulátoru se netahají.'),
        ('Verze', f'Šablona pro export předkalkulace struktury verze {d["verzeStruktury"]} (list Popis v exportu). '
                  'Vyrobeno nástrojem nastroje/sablona_podkladu_controlling.py.'),
    ]
    for r in radky:
        nav.append(list(r))
    nav['A1'].font = Font(bold=True, size=14)
    for row in nav.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = ZALOM

    # --- Skupiny ---
    hl = ['Skupina (id)', 'Název (Název budoucí analytiky)', 'Položka v ManVysl', 'Účet Helios nyní', 'Účet Helios budoucí',
          'Typ nákladu', 'Chování', 'Oblast', 'Změna', 'Poznámka controllingu']
    ws = list_s_hlavickou(wb, 'Skupiny', hl, [18, 44, 30, 18, 18, 13, 13, 9, 16, 40],
                          'Číselník nákladových skupin. Modré = dnešní stav v kalkulátoru (výchozí rozvržení z kódu). '
                          'Doplňte hlavně „Účet Helios budoucí"; Změna: ponechat / upravit / zrušit / nová.')
    for s in d['skupiny']:
        ws.append([s['id'], s['nazev'], s.get('manVysl', ''), ', '.join(s.get('ucetNyni') or []),
                   ', '.join(s.get('ucetBudouci') or []), TYP.get(s.get('typ'), ''), CHOVANI.get(s.get('chovani'), ''),
                   OBLAST.get(s.get('oblast'), ''), '', ''])
        obarvi(ws, ws.max_row, KOD_FILL)
    ws.append(['vzor-skupina', 'VZOR — Ostatní služby – příklad (smyšlené)', 'Ostatní služby k zakázkám', '999001',
               '999002, 999003', 'přímý', 'variabilní', 'OCK', 'nová', 'smyšlený řádek, smazat'])
    obarvi(ws, ws.max_row, VZOR_FILL)
    vyber(ws, 'F', ['přímý', 'nepřímý'])
    vyber(ws, 'G', ['variabilní', 'fixní'])
    vyber(ws, 'H', ['OCK', 'PROJ', 'obě'])
    vyber(ws, 'I', ['ponechat', 'upravit', 'zrušit', 'nová'])
    for row in ws.iter_rows(min_row=3, min_col=4, max_col=5):
        for c in row:
            c.number_format = '@'

    # --- Mapování PROJ ---
    hl = ['Klíč ceníku', 'Sekce ceníku PROJ', 'Položka', 'Jednotka', 'Navržená skupina (id)', 'Skupina (id) — controlling',
          'Vlastní práce / subdodávka', 'Poznámka controllingu']
    ws = list_s_hlavickou(wb, 'Mapování PROJ', hl, [24, 34, 38, 10, 20, 22, 20, 40],
                          'Každá položka ceníku projekce (PROJ) musí patřit do jedné skupiny. Navrženo podle návrhu B2 '
                          '(podklady/CONTROLLING_NAVRH_20_10.md). Potvrďte nebo přepište do sloupce „Skupina (id) — controlling".')
    for p in d['proj']:
        ws.append([p['klic'], p['sekce'], p['popis'], p['jednotka'], p['skupina'], '', '', ''])
        obarvi(ws, ws.max_row, KOD_FILL)
    ws.append(['PC.vzor', 'VZOR', 'Smyšlená položka', 'Kč/hod', 'proj-prace', 'proj-prace', 'vlastní práce', 'smyšlený řádek, smazat'])
    obarvi(ws, ws.max_row, VZOR_FILL)
    vyber(ws, 'G', ['vlastní práce', 'subdodávka', 'cestovné'])

    # --- PROJ sazby a hodiny ---
    hl = ['Role / činnost', 'Klíč ceníku (je-li)', 'Vlastní / externí', 'Nákladová sazba (Kč/h bez DPH)',
          'Co sazba obsahuje', 'Platí od', 'Platí do', 'Skupina (id)', 'Účet Helios', 'Jak se vykazují hodiny',
          'Poznámka controllingu']
    ws = list_s_hlavickou(wb, 'PROJ sazby a hodiny', hl, [26, 22, 14, 18, 34, 12, 12, 16, 14, 34, 34],
                          'Nákladové hodinové sazby pro porovnání hodin projekce (PRE = hodiny z nabídky × sazba, '
                          'POST = vykázané hodiny). Za ROLI, ne za osobu. Kalkulátor dnes zná role projektant, statik, zaměření.')
    ws.append(['VZOR — projektant (smyšlené)', 'PC.sazby.projektant', 'vlastní', 100, 'mzda + odvody (příklad)',
               '2026-01-01', '2026-12-31', 'proj-prace', '999001', 'výkaz práce po zakázkách (příklad)', 'smyšlený řádek, smazat'])
    obarvi(ws, ws.max_row, VZOR_FILL)
    ws.cell(row=ws.max_row, column=4).number_format = '#,##0.00'
    vyber(ws, 'C', ['vlastní', 'externí'])

    # --- Převodník ---
    hl = ['Číslo nabídky (CN)', 'Varianta', 'Číslo projektu Helios', 'Pipedrive deal (id)', 'Datum podpisu SoD', 'Poznámka']
    ws = list_s_hlavickou(wb, 'Převodník CN–projekt', hl, [26, 10, 22, 18, 16, 40],
                          'Jen formát: jak v Reporting CZ párujete číslo nabídky z kalkulátoru s projektem Helios '
                          '(rozhodnutí J8 — kalkulátor číslo projektu neukládá). Stačí vzorový řádek, seznam zakázek ne.')
    ws.append(['2026 - OPR - CN - 0901', 'A', 'VZOR-0001', 9001, '2026-10-14', 'smyšlený řádek'])
    obarvi(ws, ws.max_row, VZOR_FILL)

    # --- Otázky ---
    hl = ['#', 'Otázka', 'Možnosti / náš návrh', 'Odpověď controllingu']
    ws = list_s_hlavickou(wb, 'Otázky', hl, [6, 60, 60, 40],
                          'Odpovězte prosím do posledního sloupce. Návrh je to, co kalkulátor dělá dnes nebo co doporučujeme.')
    otazky = [
        ('Q1', 'Který list exportu importujete do Reporting CZ?', 'Řádky (doporučeno; jeden řádek = jedna nákladová položka) / Kapitoly (souhrn po skupinách)'),
        ('Q2', 'Formát souboru pro import', 'XLSX (dnes) / CSV — oddělovač středník, kódování UTF-8 s BOM'),
        ('Q3', 'Formát čísel v CSV', 'desetinná čárka (CZ) / desetinná tečka; bez oddělovače tisíců'),
        ('Q4', 'Podle jakého data řadit zakázky do období', 'datum prodeje = podpis SoD / objednané příplatky (dnešní výchozí) / datum odeslání nabídky'),
        ('Q5', 'Jak často a za jaké období export chcete', 'měsíčně za uplynulý měsíc / čtvrtletně / na vyžádání'),
        ('Q6', 'Jen vyhrané zakázky (Pipedrive „vyhráno")?', 'ano (dnešní výchozí) / všechny odeslané nabídky'),
        ('Q7', 'Obchodní zaokrouhlení ceny OCK nahoru (na tisíce)', 'samostatný řádek bez skupiny / rozpustit do skupin poměrem / jen v souhrnu zakázky'),
        ('Q8', 'Podíly skupin v listu Kapitoly', 'součet přesně 100 % (zbytek do největší skupiny) / zaokrouhlovat každou zvlášť (dnes)'),
        ('Q9', 'Hodiny montáže a projekce v exportu jako množství', 'ano — použijeme pro postkalkulaci osobních nákladů / ne'),
        ('Q10', 'Kdo bude mapování udržovat', 'controlling pošle vyplněnou šablonu, J. V. ji nahraje v Nastavení → Controlling'),
    ]
    for q in otazky:
        ws.append(list(q) + [''])
        for c in ws[ws.max_row]:
            c.alignment = ZALOM
    return wb


def main():
    vystup = sys.argv[1] if len(sys.argv) > 1 else VYSTUP
    os.makedirs(os.path.dirname(os.path.abspath(vystup)), exist_ok=True)
    wb = sestav(data_z_kodu())
    wb.properties.creator = 'Kalkulator Next Gen'
    wb.properties.lastModifiedBy = 'Kalkulator Next Gen'
    wb.save(vystup)
    print('OK -> ' + vystup)


if __name__ == '__main__':
    main()
